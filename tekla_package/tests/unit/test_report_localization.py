from __future__ import annotations

import csv

from openpyxl import load_workbook
from tekla_common.config.app_config import load_config, reset_config
from tekla_common.i18n import set_language
from tekla_nest.models import BarResult, NestResult, ProfileResult
from tekla_nest.services.csv_report import export_csv
from tekla_nest.services.excel_report import export_excel
from tekla_nest.services.pdf_report import render_report_html


def setup_function():
    reset_config()
    load_config("/tmp/nonexistent_config_67891.yaml")
    set_language("en")


def teardown_function():
    set_language("en")
    reset_config()


def _sample_result() -> NestResult:
    return NestResult(
        profiles=[
            ProfileResult(
                profile="IPE200",
                bars=[
                    BarResult(
                        original_length=6000,
                        mark="B1",
                        material="S355",
                        source="Market",
                        cuts=[1200, 1500],
                        cut_marks=["P1", "P2"],
                    )
                ],
                waste_pct=55.0,
                scrap_pct=5.0,
            ),
            ProfileResult(profile="HEA100", waste_pct=0.0, scrap_pct=0.0),
        ]
    )


def test_html_report_uses_current_language_for_labels():
    set_language("pt")

    html = render_report_html(_sample_result())

    assert "<title>Plano de Corte</title>" in html
    assert "Desperdício total" in html
    assert "Marca da barra" in html
    assert "Origem" in html
    assert "Não há stock disponível para este perfil." in html
    assert "Bar Mark" not in html
    assert "Overall waste" not in html


def test_csv_report_uses_current_language_for_labels(tmp_path):
    set_language("pt")

    path = export_csv(_sample_result(), tmp_path / "report.csv")

    rows = list(csv.reader(path.open(encoding="utf-8-sig"), delimiter=";"))
    assert rows[0][0] == "Plano de Corte"
    assert rows[3][0] == "Desperdício total: 55.00%"
    assert rows[5][0] == "IPE200 - Desperdício: 55.00%"
    assert rows[6][1:6] == [
        "Marca da barra",
        "Comprimento da barra",
        "Cortes",
        "Marcas dos cortes",
        "Desperdício (mm)",
    ]
    assert rows[9][0] == "HEA100 - Desperdício: 0.00%"
    assert rows[10][0] == "#"
    assert rows[11][0] == "Não há stock disponível para este perfil."


def test_excel_report_uses_current_language_for_labels(tmp_path):
    set_language("pt")

    path = export_excel(_sample_result(), tmp_path / "report.xlsx")

    workbook = load_workbook(path)
    assert "Resumo" in workbook.sheetnames
    summary = workbook["Resumo"]
    assert summary["A1"].value == "Perfil"
    assert summary["B1"].value == "Material"
    assert summary["C1"].value == "Barras usadas"
    assert summary["D1"].value == "Desperdício %"
    assert summary["E1"].value == "Sobra %"
    assert summary["A4"].value == "TOTAL"

    profile = workbook["IPE200"]
    assert profile["B1"].value == "Marca da barra"
    assert profile["C1"].value == "Comprimento da barra"
    assert profile["F1"].value == "Desperdício (mm)"
    assert profile["H1"].value == "Origem"


def test_english_report_title_replaces_default_portuguese_config_title():
    html = render_report_html(_sample_result())

    assert "<title>Cut Plan</title>" in html
    assert "Overall waste" in html
    assert "Plano de Corte" not in html
