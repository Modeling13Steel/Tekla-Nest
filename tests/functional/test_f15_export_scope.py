"""F15 — Selective export (Feedback v2 item #2.1, #6).

The user nests the whole job but on export wants to scope to one
profile/material. ``scope`` plumbs through ``render_report_html``,
``export_excel``, and ``export_csv``. Empty scope ≠ "everything";
the presenter fails the operation with a user-facing message.
"""
from __future__ import annotations

from pathlib import Path

import pytest


@pytest.fixture
def calc(journey, parts_csv_factory):
    journey.load_parts(parts_csv_factory())
    journey.auto_stock()
    journey.calculate()
    return journey


def _all_keys(result):
    from tekla_nest.services.bar_aggregation import available_scope_keys
    return available_scope_keys(result)


class TestRenderHtmlScope:
    def test_no_scope_includes_all_profiles(self, calc):
        from tekla_nest.services.pdf_report import render_report_html
        html = render_report_html(calc.presenter._last_result)
        keys = _all_keys(calc.presenter._last_result)
        for profile, _ in keys:
            assert profile in html

    def test_scope_filters_html(self, calc):
        from tekla_nest.services.pdf_report import render_report_html
        keys = _all_keys(calc.presenter._last_result)
        if len(keys) < 2:
            pytest.skip("needs ≥2 profiles to assert filtering")
        keep = keys[0]
        scope = frozenset({(keep[0].lower(), keep[1].lower())})
        html = render_report_html(calc.presenter._last_result, scope=scope)
        # Other profiles must be absent from the rendered profile-section
        # blocks. We assert at least one OTHER profile name doesn't appear
        # as a profile-section heading.
        for prof, _ in keys[1:]:
            assert f"<h2>{prof}" not in html

    def test_empty_scope_renders_empty(self, calc):
        from tekla_nest.services.pdf_report import render_report_html
        html = render_report_html(
            calc.presenter._last_result, scope=frozenset(),
        )
        assert '<div class="profile-section">' not in html


class TestPresenterEmptyScope:
    def test_export_pdf_rejects_empty_scope(self, calc, tmp_path):
        path = tmp_path / "out.pdf"
        ok = calc.presenter.export_pdf(str(path), scope=frozenset())
        assert ok is False
        assert not path.exists()

    def test_export_excel_rejects_empty_scope(self, calc, tmp_path):
        path = tmp_path / "out.xlsx"
        ok = calc.presenter.export_excel(str(path), scope=frozenset())
        assert ok is False
        assert not path.exists()

    def test_export_csv_rejects_empty_scope(self, calc, tmp_path):
        path = tmp_path / "out.csv"
        ok = calc.presenter.export_csv(str(path), scope=frozenset())
        assert ok is False
        assert not path.exists()


class TestPurchaseExportContract:
    """Feedback #6 — purchase export must carry bar length + linear m."""

    def test_excel_purchase_sheet_has_length_and_linear_m(self, calc, tmp_path):
        path = tmp_path / "out.xlsx"
        assert calc.presenter.export_excel(str(path)) is True
        from openpyxl import load_workbook
        wb = load_workbook(path)
        # Find the purchase sheet — its name comes from labels.purchase_sheet
        purchase = None
        for name in wb.sheetnames:
            if "purchas" in name.lower() or "compra" in name.lower():
                purchase = wb[name]
                break
        assert purchase is not None, "Purchase sheet missing"
        headers = [str(c.value or "").lower() for c in purchase[1]]
        joined = " ".join(headers)
        assert any("length" in h or "compriment" in h for h in headers), \
            f"bar length column missing — headers were: {headers}"
        assert "linear" in joined or "metro" in joined, \
            f"linear-metres column missing — headers were: {headers}"

    def test_csv_includes_purchase_lines(self, calc, tmp_path):
        path = tmp_path / "out.csv"
        assert calc.presenter.export_csv(str(path)) is True
        body = Path(path).read_text(encoding="utf-8-sig")
        low = body.lower()
        assert "purchas" in low or "compra" in low
        assert "linear" in low or "metro" in low
