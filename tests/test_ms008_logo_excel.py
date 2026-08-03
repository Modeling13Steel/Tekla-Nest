"""Tests for ms-008: logo visibility fix, Pillow dependency, Excel 2-sheet rework."""
from __future__ import annotations

from pathlib import Path


# ---------------------------------------------------------------------------
# TEST-001: config.yaml logo path actually exists on disk
# ---------------------------------------------------------------------------

def test_config_logo_path_exists():
    """report_logo_path resolved from config.yaml must point to an existing file."""
    from tekla_nest.config.app_config import get_config

    cfg = get_config()
    assert cfg.report_logo_path.exists(), (
        f"Logo not found at {cfg.report_logo_path!r}. "
        "Update config.yaml to reference an existing file."
    )


# ---------------------------------------------------------------------------
# Helper: build a minimal NestResult for Excel export tests
# ---------------------------------------------------------------------------

def _minimal_result():
    from tekla_nest.models.nest_result import NestResult, ProfileResult
    from tekla_nest.models.bar_result import BarResult

    bar = BarResult(
        original_length=6100.0,
        mark="B1",
        material="S275JR",
        source="stock",
        cuts=[3000.0, 2500.0],
        cut_marks=["P1", "P2"],
    )
    prof = ProfileResult(
        profile="HEA240",
        bars=[bar],
        waste_pct=9.8,
        scrap_pct=0.0,
        unfit_pieces=[],
        material="S275JR",
    )
    return NestResult(profiles=[prof])


# ---------------------------------------------------------------------------
# TEST-002: exactly 2 sheets
# ---------------------------------------------------------------------------

def test_excel_two_sheets_only(tmp_path):
    """export_excel() must produce a workbook with exactly 2 sheets."""
    import openpyxl
    from tekla_nest.services.excel_report import export_excel

    out = tmp_path / "test_ms008.xlsx"
    export_excel(_minimal_result(), str(out))

    wb = openpyxl.load_workbook(str(out))
    assert len(wb.sheetnames) == 2, (
        f"Expected 2 sheets, got {len(wb.sheetnames)}: {wb.sheetnames}"
    )


# ---------------------------------------------------------------------------
# TEST-003: sheet names match Summary / Purchase (or Portuguese equivalents)
# ---------------------------------------------------------------------------

def test_excel_sheet_names(tmp_path):
    """First sheet should be Summary (or Portuguese), second Purchase."""
    import openpyxl
    from tekla_nest.services.excel_report import export_excel

    out = tmp_path / "test_ms008_names.xlsx"
    export_excel(_minimal_result(), str(out))

    wb = openpyxl.load_workbook(str(out))
    names = wb.sheetnames

    # Accept English or Portuguese variants
    summary_variants = ("summary", "resumo", "sumário", "sumario")
    purchase_variants = ("purchase", "compra", "encomenda", "aquisição", "aquisicao")

    assert names[0].lower().startswith(summary_variants), (
        f"Sheet 0 name {names[0]!r} is not a Summary variant"
    )
    assert names[1].lower().startswith(purchase_variants), (
        f"Sheet 1 name {names[1]!r} is not a Purchase variant"
    )


# ---------------------------------------------------------------------------
# TEST-004: Summary sheet contains the profile name
# ---------------------------------------------------------------------------

def test_excel_summary_has_profile_data(tmp_path):
    """Summary sheet must contain the profile name somewhere in its cell values."""
    import openpyxl
    from tekla_nest.services.excel_report import export_excel

    out = tmp_path / "test_ms008_content.xlsx"
    result = _minimal_result()
    export_excel(result, str(out))

    wb = openpyxl.load_workbook(str(out))
    ws = wb.worksheets[0]  # Summary sheet

    all_values = [
        str(cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    ]
    profile_name = result.profiles[0].profile  # "HEA240"
    assert any(profile_name in v for v in all_values), (
        f"Profile name {profile_name!r} not found in Summary sheet. "
        f"Values found: {all_values[:20]}"
    )


# ---------------------------------------------------------------------------
# TEST-005: Pillow appears in pyproject.toml dependencies
# ---------------------------------------------------------------------------

def test_pillow_in_dependencies():
    """pyproject.toml must list Pillow as a dependency."""
    toml_path = Path(__file__).parent.parent / "pyproject.toml"
    content = toml_path.read_text(encoding="utf-8")
    assert "Pillow" in content, (
        "Pillow not found in pyproject.toml. "
        "Add 'Pillow>=9.0' to the dependencies array."
    )
