"""Feedback #10 — Excel + CSV exports must carry a purchase aggregation
suitable for ordering: profile × material × bar length × count, with
linear-metre subtotals per (profile, material) and a grand total.
"""
from __future__ import annotations

import csv

from openpyxl import load_workbook

from tekla_nest.config.app_config import load_config, reset_config
from tekla_nest.i18n import set_language
from tekla_nest.models import BarResult, NestResult, ProfileResult
from tekla_nest.services.csv_report import export_csv
from tekla_nest.services.excel_report import export_excel


def setup_function():
    reset_config()
    load_config("/tmp/nonexistent_config_purchase.yaml")
    set_language("en")


def teardown_function():
    set_language("en")
    reset_config()


def _multi_material_result() -> NestResult:
    """Two profiles, HEA240 split across two materials and two lengths."""
    return NestResult(
        profiles=[
            ProfileResult(
                profile="HEA240",
                material="S235JR",
                bars=[
                    BarResult(original_length=6000, mark="HEA240",
                              material="S235JR", source="Mercado",
                              cuts=[3000], cut_marks=["A1"]),
                    BarResult(original_length=6000, mark="HEA240",
                              material="S235JR", source="Mercado",
                              cuts=[3000], cut_marks=["A2"]),
                ],
                waste_pct=50.0,
            ),
            ProfileResult(
                profile="HEA240",
                material="S275JR",
                bars=[
                    BarResult(original_length=12000, mark="HEA240",
                              material="S275JR", source="Mercado",
                              cuts=[5000], cut_marks=["B1"]),
                ],
                waste_pct=58.3,
            ),
            ProfileResult(
                profile="IPE200",
                material="S355JR",
                bars=[
                    BarResult(original_length=6000, mark="IPE200",
                              material="S355JR", source="Mercado",
                              cuts=[2800, 2800], cut_marks=["C1", "C2"]),
                ],
                waste_pct=6.7,
            ),
        ]
    )


# ── Excel ────────────────────────────────────────────────────────


def test_excel_has_purchase_sheet(tmp_path):
    path = export_excel(_multi_material_result(), tmp_path / "report.xlsx")
    wb = load_workbook(path)
    assert "Purchase" in wb.sheetnames


def test_excel_purchase_sheet_aggregates_by_profile_material_length(tmp_path):
    path = export_excel(_multi_material_result(), tmp_path / "report.xlsx")
    wb = load_workbook(path)
    ws = wb["Purchase"]

    # Headers
    assert ws["A1"].value == "Profile"
    assert ws["B1"].value == "Material"
    assert ws["C1"].value == "Bar Length"
    assert ws["D1"].value == "Bars Used"
    assert ws["E1"].value == "Linear m"

    # Collect data rows (skip subtotal rows that have empty C column)
    rows = []
    for r in range(2, ws.max_row + 1):
        a, b, c, d, e = (
            ws.cell(row=r, column=col).value for col in range(1, 6)
        )
        if c is None or c == "":
            continue
        rows.append((a, b, c, d, e))

    # Expect 3 distinct (profile, material, length) entries.
    assert ("HEA240", "S235JR", 6000, 2, 12.0) in rows
    assert ("HEA240", "S275JR", 12000, 1, 12.0) in rows
    assert ("IPE200", "S355JR", 6000, 1, 6.0) in rows
    assert len(rows) == 3


def test_excel_purchase_sheet_includes_per_profile_subtotals(tmp_path):
    path = export_excel(_multi_material_result(), tmp_path / "report.xlsx")
    wb = load_workbook(path)
    ws = wb["Purchase"]

    # Scan column A for subtotal labels.
    a_col = [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]
    labels = [v for v in a_col if isinstance(v, str)]

    assert any("HEA240 / S235JR" in v for v in labels)
    assert any("HEA240 / S275JR" in v for v in labels)
    assert any("IPE200 / S355JR" in v for v in labels)


def test_excel_purchase_sheet_grand_total_matches_sum_of_bars(tmp_path):
    path = export_excel(_multi_material_result(), tmp_path / "report.xlsx")
    wb = load_workbook(path)
    ws = wb["Purchase"]

    # Grand total = 2*6 + 1*12 + 1*6 = 30 linear m.
    last_row = ws.max_row
    assert ws.cell(row=last_row, column=1).value == "TOTAL"
    assert abs(float(ws.cell(row=last_row, column=5).value) - 30.0) < 0.01


# ── CSV ──────────────────────────────────────────────────────────


def test_csv_contains_purchase_section(tmp_path):
    path = export_csv(_multi_material_result(), tmp_path / "report.csv")
    text = path.read_text(encoding="utf-8-sig")
    assert "Purchase" in text
    # Grand total appears at the bottom.
    assert "30.00" in text


def test_csv_purchase_section_has_correct_aggregation(tmp_path):
    path = export_csv(_multi_material_result(), tmp_path / "report.csv")
    rows = list(csv.reader(path.open(encoding="utf-8-sig"), delimiter=";"))

    # Find the Purchase header row.
    purchase_idx = next(
        i for i, r in enumerate(rows) if r and r[0] == "Purchase"
    )
    # Header is one row down.
    header = rows[purchase_idx + 1]
    assert header[:5] == ["Profile", "Material", "Bar Length",
                          "Bars Used", "Linear m"]

    data_rows = []
    for r in rows[purchase_idx + 2:]:
        if not r or not r[0]:
            continue
        if r[0].startswith("TOTAL"):
            continue
        if len(r) >= 5 and r[2] and r[3]:
            data_rows.append((r[0], r[1], int(r[2]), int(r[3]), float(r[4])))

    assert ("HEA240", "S235JR", 6000, 2, 12.0) in data_rows
    assert ("HEA240", "S275JR", 12000, 1, 12.0) in data_rows
    assert ("IPE200", "S355JR", 6000, 1, 6.0) in data_rows


def test_excel_per_material_sheets_do_not_collide(tmp_path):
    """Feedback #5 + #10 — same profile name with two materials must
    yield two distinct sheets, not silently overwrite."""
    path = export_excel(_multi_material_result(), tmp_path / "report.xlsx")
    wb = load_workbook(path)
    names = wb.sheetnames
    # Both HEA240 materials should be present as separate sheets.
    hea_sheets = [n for n in names if n.startswith("HEA240")]
    assert len(hea_sheets) == 2
