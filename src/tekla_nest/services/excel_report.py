"""Excel report generation — Summary sheet + Purchase sheet.

Requires ``openpyxl``.  Install with:  pip install openpyxl
"""
from __future__ import annotations

import contextlib
import logging
from datetime import datetime
from pathlib import Path

from ..config.app_config import get_config
from ..models import NestResult
from .report_labels import report_labels

LOGGER = logging.getLogger(__name__)


def export_excel(
    result: NestResult,
    output_path: str | Path,
    scope: frozenset[tuple[str, str]] | None = None,
    attached_image_path: str | Path | None = None,
    project_name: str = "",
    milestone: str = "",
) -> Path:
    """Write nesting results to an Excel workbook.

    Produces exactly 2 sheets:
        1. Summary — full PDF mirror: logo, KPI block, per-profile sections
           with data/strip rows, unfit pieces, and optional site image.
        2. Purchase — aggregation suitable for ordering
           (profile × material × bar length × count + linear-metre subtotals).

    Returns the resolved output path.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.drawing.image import Image as XlImage
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        raise RuntimeError(
            "openpyxl is not installed.  "
            "Install with:  pip install openpyxl"
        )

    cfg = get_config()
    labels = report_labels(cfg)
    from .bar_aggregation import aggregate_prep, filter_result  # noqa: PLC0415
    result = filter_result(result, scope)
    wb = Workbook()

    # ── Summary sheet ─────────────────────────────────────────
    ws_summary = wb.active
    ws_summary.title = labels.summary_sheet[:31]

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor=cfg.primary_color.lstrip("#"))
    wrap = Alignment(wrap_text=True, vertical="center")

    # N_COLS used for strip rows and merges
    N_COLS = 6 + (1 if cfg.show_material else 0) + (1 if cfg.show_stock_source else 0)

    # Logo — placed at A1, data starts at row 5 to leave room
    LOGO_ROW_OFFSET = 4  # rows reserved for logo
    logo_placed = False
    try:
        from ..design_system.brand import select_logo_variant
        logo_path = select_logo_variant(cfg.report_logo_path, "#FFFFFF")
        if logo_path.exists() and logo_path.suffix.lower() in (".png", ".jpg", ".jpeg", ".bmp", ".gif"):
            img = XlImage(str(logo_path))
            # Scale to 50px height, preserve aspect ratio
            orig_h = img.height or 1
            orig_w = img.width or 150
            img.height = 50
            img.width = int(orig_w * (50 / orig_h))
            ws_summary.add_image(img, "A1")
            logo_placed = True
    except ImportError:
        LOGGER.warning("Pillow not installed; Excel logo will be skipped. Run: pip install Pillow>=9.0")
    except Exception as exc:
        LOGGER.warning("Logo could not be embedded in Excel: %s", exc)

    data_start_row = (LOGO_ROW_OFFSET + 1) if logo_placed else 1
    current_row = data_start_row

    # ── Company name + timestamp ──────────────────────────────
    ws_summary.cell(row=current_row, column=1, value=cfg.company_name).font = Font(bold=True)
    ws_summary.cell(
        row=current_row, column=2, value=datetime.now().strftime("%Y-%m-%d %H:%M")
    )
    current_row += 1

    if project_name:
        ws_summary.cell(
            row=current_row, column=1,
            value=f"{labels.project_name_label}: {project_name}",
        ).font = Font(bold=True)
        current_row += 1
    if milestone:
        ws_summary.cell(
            row=current_row, column=1,
            value=f"{labels.milestone_label}: {milestone}",
        ).font = Font(italic=True)
        current_row += 1

    # Empty row
    current_row += 1

    # ── KPI block (rows 4–7 relative to data_start_row) ──────
    total_bars = sum(len(p.bars) for p in result.profiles)
    unfit_count = sum(len(p.unfit_pieces) for p in result.profiles)
    profile_count = len(result.profiles)

    kpi_rows = [
        ("Overall Waste %", round(result.overall_waste_pct, 2)),
        ("Total Bars", total_bars),
        ("Unfit Pieces", unfit_count),
        ("Profiles", profile_count),
    ]
    for label_text, value in kpi_rows:
        ws_summary.cell(row=current_row, column=1, value=label_text).font = Font(bold=True)
        ws_summary.cell(row=current_row, column=2, value=value)
        current_row += 1

    # Empty row
    current_row += 1

    # ── Per-profile sections ──────────────────────────────────
    preps = aggregate_prep(result)
    prep_by_key = {
        (p.profile.lower().strip(), (p.material or "").lower().strip()): p
        for p in preps
    }

    WASTE_FILL = PatternFill("solid", fgColor="CCCCCC")
    CUT_COLORS = ["4472C4", "ED7D31", "A9D18E", "FF0000", "70AD47", "FFC000"]

    columns = [
        "#",
        labels.bar_mark,
        labels.bar_length,
        labels.cuts,
        labels.cut_marks,
        labels.waste_mm,
    ]
    if cfg.show_material:
        columns.append(labels.material)
    if cfg.show_stock_source:
        columns.append(labels.source)

    for prof in result.profiles:
        # a. Profile header row — merged, bold white on primary fill
        prof_header_text = f"{prof.profile} — {prof.material}" if prof.material else prof.profile
        if N_COLS > 1:
            ws_summary.merge_cells(
                start_row=current_row, start_column=1,
                end_row=current_row, end_column=N_COLS,
            )
        cell = ws_summary.cell(row=current_row, column=1, value=prof_header_text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=False, vertical="center")
        current_row += 1

        # b. Prep summary row — merged
        prof_key = (prof.profile.lower().strip(), (prof.material or "").lower().strip())
        prep = prep_by_key.get(prof_key)
        if prep:
            prep_parts = ", ".join(f"{cnt}× {length_mm:.0f}mm" for length_mm, cnt in prep.by_length)
            prep_text = f"Prep: {prep_parts} | {prep.total_bars} bars · {prep.linear_m:.1f} m"
            if N_COLS > 1:
                ws_summary.merge_cells(
                    start_row=current_row, start_column=1,
                    end_row=current_row, end_column=N_COLS,
                )
            cell = ws_summary.cell(row=current_row, column=1, value=prep_text)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=False, vertical="center")
        current_row += 1

        # c. Column header row
        for col, h in enumerate(columns, 1):
            cell = ws_summary.cell(row=current_row, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = wrap
        current_row += 1

        # d. Data rows + strip rows
        for bar_idx, bar in enumerate(prof.bars):
            # Data row
            ws_summary.cell(row=current_row, column=1, value=bar_idx + 1)
            ws_summary.cell(row=current_row, column=2, value=bar.mark)
            ws_summary.cell(row=current_row, column=3, value=bar.original_length)
            ws_summary.cell(row=current_row, column=4, value=", ".join(str(int(c)) for c in bar.cuts))
            ws_summary.cell(row=current_row, column=5, value=", ".join(bar.cut_marks))
            ws_summary.cell(row=current_row, column=6, value=round(bar.free_length, 0))
            col_offset = 7
            if cfg.show_material:
                ws_summary.cell(row=current_row, column=col_offset, value=bar.material)
                col_offset += 1
            if cfg.show_stock_source:
                ws_summary.cell(row=current_row, column=col_offset, value=bar.source)

            # Strip row (visual bar)
            strip_row = current_row + 1
            ws_summary.row_dimensions[strip_row].height = 12

            bar_total = bar.original_length or 1
            col_start = 1
            for i, cut in enumerate(bar.cuts):
                span = max(1, round(cut / bar_total * N_COLS))
                remaining = N_COLS - col_start + 1
                if remaining <= 0:
                    break
                span = min(span, remaining)
                col_end = col_start + span - 1
                if col_end > col_start:
                    with contextlib.suppress(Exception):
                        ws_summary.merge_cells(
                            start_row=strip_row, start_column=col_start,
                            end_row=strip_row, end_column=col_end,
                        )
                color = CUT_COLORS[i % len(CUT_COLORS)]
                ws_summary.cell(row=strip_row, column=col_start).fill = PatternFill(
                    "solid", fgColor=color
                )
                col_start += span

            # Remaining columns = waste (grey)
            if col_start <= N_COLS:
                col_end = N_COLS
                if col_end > col_start:
                    with contextlib.suppress(Exception):
                        ws_summary.merge_cells(
                            start_row=strip_row, start_column=col_start,
                            end_row=strip_row, end_column=col_end,
                        )
                ws_summary.cell(row=strip_row, column=col_start).fill = WASTE_FILL

            current_row += 2  # data row + strip row

        # e. Unfit pieces
        if prof.unfit_pieces:
            ws_summary.cell(row=current_row, column=1, value="Unfit Pieces:").font = Font(bold=True)
            current_row += 1
            for piece in prof.unfit_pieces:
                ws_summary.cell(row=current_row, column=1, value=piece.mark)
                ws_summary.cell(row=current_row, column=2, value=piece.length)
                current_row += 1

        # f. Blank separator row
        current_row += 1

    # ── Site image embedded at end of Summary sheet ───────────
    if attached_image_path:
        img_path = Path(attached_image_path)
        if img_path.exists():
            try:
                site_img = XlImage(str(img_path))
                site_img.height = 400
                site_img.width = 600
                ws_summary.add_image(site_img, f"A{current_row + 1}")
            except Exception as exc:
                LOGGER.warning("Site image could not be embedded in Excel: %s", exc)
                ws_summary.cell(row=current_row + 1, column=1, value=str(img_path))

    _auto_column_width(ws_summary)

    # ── Purchase sheet (feedback #10) ─────────────────────────
    _write_purchase_sheet(wb, result, cfg, labels, header_font, header_fill, wrap)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(str(out))
    except PermissionError:
        raise RuntimeError(
            f"Cannot save '{out.name}': the file is open in another application.\n"
            f"  Fix: Close '{out.name}' in Excel (or any other program) and export again."
        )
    return out


def _write_purchase_sheet(wb, result, cfg, labels,
                          header_font, header_fill, wrap) -> None:
    """Build a purchase-order-friendly aggregation sheet.

    Rows: (profile, material, bar length) × quantity, plus a per-profile
    linear-metre subtotal row. Feedback #10 — the user wants a single
    glance-able sheet they can hand to a supplier.
    """
    from openpyxl.styles import Alignment, Font

    ws = wb.create_sheet(title=labels.purchase_sheet[:31])

    headers = [
        labels.profile,
        labels.material,
        labels.bar_length,
        labels.bars_used,
        labels.linear_meters,
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = wrap

    # Aggregate: {(profile, material, length): count}
    agg: dict[tuple[str, str, float], int] = {}
    for prof in result.profiles:
        for bar in prof.bars:
            key = (prof.profile, prof.material or bar.material,
                   float(bar.original_length))
            agg[key] = agg.get(key, 0) + 1

    if not agg:
        ws.cell(row=2, column=1, value=labels.no_stock)
        _auto_column_width(ws)
        return

    # Sort by profile, material, length DESC
    sorted_keys = sorted(agg.keys(), key=lambda k: (k[0], k[1], -k[2]))

    row = 2
    bold = Font(bold=True)
    grand_total_m = 0.0
    last_group: tuple[str, str] | None = None
    group_total_m = 0.0

    def _flush_subtotal(profile: str, material: str, total_m: float, r: int) -> int:
        label = f"{profile}" + (f" / {material}" if material else "")
        cell = ws.cell(row=r, column=1, value=f"{labels.total} {label}")
        cell.font = bold
        cell.alignment = Alignment(horizontal="right")
        ws.cell(row=r, column=5, value=round(total_m, 2)).font = bold
        return r + 1

    for key in sorted_keys:
        profile, material, length = key
        count = agg[key]
        linear_m = (length * count) / 1000.0

        if last_group is not None and (profile, material) != last_group:
            row = _flush_subtotal(last_group[0], last_group[1], group_total_m, row)
            group_total_m = 0.0

        ws.cell(row=row, column=1, value=profile)
        ws.cell(row=row, column=2, value=material)
        ws.cell(row=row, column=3, value=length)
        ws.cell(row=row, column=4, value=count)
        ws.cell(row=row, column=5, value=round(linear_m, 2))
        row += 1

        group_total_m += linear_m
        grand_total_m += linear_m
        last_group = (profile, material)

    if last_group is not None:
        row = _flush_subtotal(last_group[0], last_group[1], group_total_m, row)

    cell = ws.cell(row=row, column=1, value=labels.total)
    cell.font = bold
    cell.alignment = Alignment(horizontal="right")
    ws.cell(row=row, column=5, value=round(grand_total_m, 2)).font = bold

    _auto_column_width(ws)


def _auto_column_width(ws) -> None:
    """Set column widths based on content length.

    Skips MergedCell objects (they don't have column_letter / value to
    contribute meaningful width data — the merge anchor handles it).
    """
    from openpyxl.cell.cell import MergedCell
    for col in ws.columns:
        max_len = 0
        col_letter = None
        for cell in col:
            if isinstance(cell, MergedCell):
                continue
            if col_letter is None:
                col_letter = cell.column_letter
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        if col_letter is not None:
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)
